#!/usr/bin/env python3
"""Build the first-exam report and its UML figures from the project Markdown."""

from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "informe-primer-examen.md"
OUTPUT = ROOT / "docs" / "entregables" / "informe-primer-examen.docx"
FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
NAVY = "17324D"
BLUE = "245B84"
PALE_BLUE = "EAF1F7"
INK = "1D2730"
MUTED = "4C5B67"
LINE = "D4DCE3"
PAPER = "F8FAFC"


def rgb(value: str) -> tuple[int, int, int]:
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size)


def wrap(draw: ImageDraw.ImageDraw, text: str, typeface: ImageFont.FreeTypeFont, width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if current and draw.textbbox((0, 0), candidate, font=typeface)[2] > width:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    text: str,
    xy: tuple[int, int],
    typeface: ImageFont.FreeTypeFont,
    fill: str,
    width: int,
    line_gap: int = 8,
) -> int:
    x, y = xy
    for line in wrap(draw, text, typeface, width):
        draw.text((x, y), line, font=typeface, fill=f"#{fill}")
        y += typeface.size + line_gap
    return y


def arrowhead(draw: ImageDraw.ImageDraw, end: tuple[int, int], before: tuple[int, int], color: str) -> None:
    x, y = end
    bx, by = before
    size = 20
    if abs(x - bx) >= abs(y - by):
        sign = 1 if x > bx else -1
        points = [(x, y), (x - sign * size, y - size // 2), (x - sign * size, y + size // 2)]
    else:
        sign = 1 if y > by else -1
        points = [(x, y), (x - size // 2, y - sign * size), (x + size // 2, y - sign * size)]
    draw.polygon(points, fill=f"#{color}")


def route(
    draw: ImageDraw.ImageDraw,
    points: list[tuple[int, int]],
    label: str,
    label_at: tuple[int, int],
    start_mult: str = "",
    end_mult: str = "",
    dashed: bool = False,
) -> None:
    color = BLUE
    for start, end in zip(points, points[1:]):
        if dashed:
            x1, y1 = start
            x2, y2 = end
            if y1 == y2:
                lo, hi = sorted((x1, x2))
                for x in range(lo, hi, 38):
                    draw.line((x, y1, min(x + 20, hi), y1), fill=f"#{color}", width=5)
            else:
                lo, hi = sorted((y1, y2))
                for y in range(lo, hi, 38):
                    draw.line((x1, y, x1, min(y + 20, hi)), fill=f"#{color}", width=5)
        else:
            draw.line((start, end), fill=f"#{color}", width=5, joint="curve")
    arrowhead(draw, points[-1], points[-2], color)
    if label:
        label_font = font(32, True)
        text = label
        bounds = draw.textbbox((0, 0), text, font=label_font)
        pad_x, pad_y = 12, 8
        x, y = label_at
        draw.rounded_rectangle(
            (x - pad_x, y - pad_y, x + bounds[2] + pad_x, y + bounds[3] + pad_y),
            radius=12,
            fill=f"#{PAPER}",
            outline=f"#{LINE}",
            width=2,
        )
        draw.text((x, y), text, font=label_font, fill=f"#{NAVY}")
    mult_font = font(28, True)
    if start_mult:
        x1, y1 = points[0]
        x2, y2 = points[1]
        sx = 1 if x2 > x1 else -1 if x2 < x1 else 0
        sy = 1 if y2 > y1 else -1 if y2 < y1 else 0
        draw.text((x1 + sx * 15 + 5, y1 + sy * 15 + 5), start_mult, font=mult_font, fill=f"#{INK}")
    if end_mult:
        x1, y1 = points[-2]
        x2, y2 = points[-1]
        sx = 1 if x2 > x1 else -1 if x2 < x1 else 0
        sy = 1 if y2 > y1 else -1 if y2 < y1 else 0
        segment_length = abs(x2 - x1) + abs(y2 - y1)
        offset = 35 if segment_length < 160 else 70
        draw.text((x2 - sx * offset + 5, y2 - sy * offset + 5), end_mult, font=mult_font, fill=f"#{INK}")


def class_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    attributes: list[str],
    operations: list[str] | None = None,
) -> None:
    x, y, w, h = box
    draw.rounded_rectangle((x, y, x + w, y + h), radius=20, fill="#FFFFFF", outline=f"#{LINE}", width=4)
    head_h = 92
    draw.rounded_rectangle((x, y, x + w, y + head_h + 20), radius=20, fill=f"#{NAVY}")
    draw.rectangle((x, y + head_h, x + w, y + head_h + 15), fill=f"#{NAVY}")
    title_font = font(44, True)
    title_width = draw.textbbox((0, 0), title, font=title_font)[2]
    draw.text((x + (w - title_width) / 2, y + 17), title, font=title_font, fill="#FFFFFF")
    y_text = y + head_h + 30
    attr_font = font(32)
    for item in attributes:
        y_text = draw_wrapped(draw, item, (x + 28, y_text), attr_font, INK, w - 56, line_gap=2) + 5
    if operations:
        draw.line((x + 22, y_text + 8, x + w - 22, y_text + 8), fill=f"#{LINE}", width=3)
        y_text += 22
        for item in operations:
            y_text = draw_wrapped(draw, item, (x + 28, y_text), attr_font, BLUE, w - 56, line_gap=2) + 3


def figure_header(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    title_font = font(62, True)
    subtitle_font = font(32)
    draw.text((90, 42), title, font=title_font, fill=f"#{NAVY}")
    draw.text((92, 120), subtitle, font=subtitle_font, fill=f"#{MUTED}")


def draw_domain(path: Path) -> None:
    image = Image.new("RGB", (3200, 2000), f"#{PAPER}")
    draw = ImageDraw.Draw(image)
    figure_header(draw, "Modelo de Dominio", "Clases y asociaciones presentes en el esquema SQL proporcionado")

    # Associations are drawn before the class boxes so lines stay behind class borders.
    route(draw, [(430, 590), (430, 980)], "cuenta", (455, 770), "1", "0..1")
    route(draw, [(780, 390), (1000, 390)], "solicitante", (800, 325), "1", "0..*")
    route(draw, [(2570, 430), (2570, 600)], "", (0, 0), "1", "0..*")
    route(draw, [(2570, 860), (2570, 1030)], "", (0, 0), "1", "0..*")
    route(draw, [(2570, 1290), (2570, 1460)], "", (0, 0), "1", "0..*")
    route(draw, [(780, 540), (850, 650), (850, 1160), (970, 1160)], "persona", (855, 850), "1", "0..*")
    route(draw, [(1500, 660), (1500, 820), (900, 820), (900, 1040), (970, 1040)], "reserva origen", (1120, 790), "0..1", "0..1")
    route(draw, [(1670, 530), (1900, 530), (1900, 1500), (2140, 1500)], "unidad", (1735, 480), "0..*", "1")
    route(draw, [(1670, 1170), (2020, 1170), (2020, 1570), (2140, 1570)], "unidad", (1760, 1120), "0..*", "1")
    route(draw, [(1320, 1480), (1320, 1590)], "renovaciones", (1410, 1510), "1", "0..*")

    class_box(draw, (80, 170, 700, 420), "Persona", [
        "idPersona: UUID", "cui, documento, nombre", "correoInstitucional: String", "rolBase: Rol", "vinculacion, estado",
    ])
    class_box(draw, (80, 980, 700, 350), "Cuenta", [
        "idCuenta: UUID", "nombreUsuario: String", "hashContrasena: String", "estado, ultimoAcceso",
    ])
    class_box(draw, (1000, 170, 670, 490), "Reserva", [
        "idReserva: UUID", "intervaloInicio, intervaloFin", "fechaSolicitud, usoPrevisto", "estado, motivoRechazo",
    ], ["confirmar()  rechazar()  cancelar()"])
    class_box(draw, (2140, 170, 860, 260), "CategoriaBien", ["idCategoria: UUID", "nombre, descripcion, estado"])
    class_box(draw, (2140, 600, 860, 260), "TipoBien", ["idTipo: UUID", "nombre, esPrestable, estado"])
    class_box(draw, (2140, 1030, 860, 260), "FichaBien", ["idFicha: UUID", "nombre, descripcion, estado"])
    class_box(draw, (2140, 1460, 860, 400), "UnidadFisica", [
        "idUnidad: UUID", "codigoInventario: String", "numeroSerie, ubicacion", "condicionFisica, accesorios", "estado: EstadoUnidad",
    ], ["estaDisponible()"])
    class_box(draw, (970, 980, 700, 500), "Prestamo", [
        "idPrestamo: UUID", "plazoInicio, plazoVencimiento", "estado, modalidad", "datos de entrega y devolucion",
    ], ["registrarEntrega()  registrarDevolucion()"])
    class_box(draw, (970, 1590, 700, 300), "Renovacion", [
        "idRenovacion: UUID", "secuencia: Integer", "vencimientos anterior y nuevo", "administrador, fechaAutorizacion",
    ])

    draw.text((90, 1945), "Persona participa como solicitante y administrador; las relaciones siguen las claves foraneas del SQL.", font=font(27), fill=f"#{MUTED}")
    image.save(path, dpi=(300, 300))


def package_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    packages: list[tuple[str, list[str]]],
) -> None:
    x, y, w, h = box
    draw.rounded_rectangle((x, y, x + w, y + h), radius=24, fill="#FFFFFF", outline=f"#{LINE}", width=4)
    title_font = font(46, True)
    draw.rounded_rectangle((x, y, x + w, y + 88), radius=24, fill=f"#{NAVY}")
    draw.rectangle((x, y + 60, x + w, y + 90), fill=f"#{NAVY}")
    draw.text((x + 25, y + 15), title, font=title_font, fill="#FFFFFF")
    current_y = y + 120
    for package_name, classes in packages:
        draw.text((x + 26, current_y), f"<<package>> {package_name}", font=font(30, True), fill=f"#{BLUE}")
        current_y += 48
        for item in classes:
            current_y = draw_wrapped(draw, item, (x + 44, current_y), font(29), INK, w - 78, line_gap=1) + 6
        current_y += 22


def draw_architecture(path: Path) -> None:
    image = Image.new("RGB", (3200, 2000), f"#{PAPER}")
    draw = ImageDraw.Draw(image)
    figure_header(draw, "Arquitectura en Capas", "Paquetes, clases representativas y direccion de dependencias")
    boxes = [
        (70, 210, 700, 1390),
        (850, 210, 750, 1390),
        (1680, 210, 760, 1390),
        (2520, 210, 610, 1390),
    ]
    # Layer-to-layer dependencies point inward toward application contracts and domain rules.
    route(draw, [(770, 830), (840, 830)], "usa", (770, 775))
    route(draw, [(1600, 830), (1670, 830)], "usa", (1605, 775))
    route(draw, [(2520, 1120), (2460, 1120)], "", (0, 0), dashed=True)
    route(draw, [(2825, 1600), (2825, 1740)], "SQL", (2845, 1645))

    package_box(draw, boxes[0], "Presentacion", [
        ("api.auth", ["AuthRouter", "TokenSchema"]),
        ("api.inventory", ["InventoryRouter", "CreateUnitSchema"]),
        ("api.reservations", ["ReservationRouter"]),
        ("api.loans", ["LoanRouter", "ReturnSchema"]),
    ])
    package_box(draw, boxes[1], "Aplicacion", [
        ("use_cases", ["AuthenticateUser", "CreateUnit", "CreateReservation", "RegisterLoan", "RecordReturn"]),
        ("ports", ["UnitRepository", "LoanRepository", "UnitOfWork", "PasswordHasher"]),
        ("dto", ["Commands and results"]),
    ])
    package_box(draw, boxes[2], "Dominio", [
        ("identity", ["Persona", "Cuenta", "Rol"]),
        ("inventory", ["CategoriaBien", "TipoBien", "FichaBien", "UnidadFisica"]),
        ("reservations", ["Reserva", "EstadoReserva"]),
        ("loans", ["Prestamo", "Renovacion", "EstadoPrestamo"]),
        ("contracts", ["Repository interfaces"]),
    ])
    package_box(draw, boxes[3], "Infraestructura", [
        ("persistence", ["SQLAlchemy mappings", "Postgres repositories", "Session"]),
        ("security", ["Argon2PasswordHasher", "JwtTokenService"]),
        ("migrations", ["Alembic"]),
        ("composition", ["Dependency wiring"]),
    ])
    # Database cylinder.
    x1, y1, x2, y2 = 2590, 1760, 3090, 1950
    draw.ellipse((x1, y1, x2, y1 + 65), fill="#EAF1F7", outline=f"#{BLUE}", width=5)
    draw.rectangle((x1, y1 + 32, x2, y2 - 32), fill="#EAF1F7", outline=f"#{BLUE}", width=5)
    draw.arc((x1, y2 - 65, x2, y2 + 5), 0, 180, fill=f"#{BLUE}", width=5)
    db_text = "PostgreSQL"
    text_width = draw.textbbox((0, 0), db_text, font=font(40, True))[2]
    draw.text(((x1 + x2 - text_width) / 2, y1 + 70), db_text, font=font(40, True), fill=f"#{NAVY}")
    draw.text((90, 1910), "Persistencia implementa contratos del dominio; los casos de uso coordinan las reglas y operaciones.", font=font(27), fill=f"#{MUTED}")
    image.save(path, dpi=(300, 300))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top: int = 110, start: int = 110, bottom: int = 110, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "5")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), LINE)
        borders.append(border)
    tbl_pr.append(borders)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def set_run_font(run, *, name: str = "Arial", size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\[[^\]]+\]\([^)]+\))", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Menlo", size=9, color=BLUE)
        elif part.startswith("[") and "](" in part:
            label, url = part[1:].split("](", 1)
            url = url.rstrip(")")
            run = paragraph.add_run(label)
            set_run_font(run, color=BLUE)
            run.underline = True
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal_rfonts = normal._element.rPr.rFonts
    if normal_rfonts is None:
        normal_rfonts = OxmlElement("w:rFonts")
        normal._element.rPr.insert(0, normal_rfonts)
    normal_rfonts.set(qn("w:ascii"), "Arial")
    normal_rfonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Arial"
        rfonts = style._element.rPr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            style._element.rPr.insert(0, rfonts)
        rfonts.set(qn("w:ascii"), "Arial")
        rfonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Arial"
    title_rfonts = title._element.rPr.rFonts
    if title_rfonts is None:
        title_rfonts = OxmlElement("w:rFonts")
        title._element.rPr.insert(0, title_rfonts)
    title_rfonts.set(qn("w:ascii"), "Arial")
    title_rfonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    if "Caption" in styles:
        styles["Caption"].font.name = "Arial"
        rfonts = styles["Caption"]._element.rPr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            styles["Caption"]._element.rPr.insert(0, rfonts)
        rfonts.set(qn("w:ascii"), "Arial")
        rfonts.set(qn("w:hAnsi"), "Arial")
        styles["Caption"].font.size = Pt(9)
        styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)
        styles["Caption"].paragraph_format.space_before = Pt(7)
        styles["Caption"].paragraph_format.space_after = Pt(4)

    doc.core_properties.title = "Sistema de gestión de préstamos de bienes de la EPCC"
    doc.core_properties.subject = "Informe para el primer examen"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "EPCC, requisitos, modelo de dominio, UML, arquitectura en capas"
    return doc


def landscape_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    return section


def portrait_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    return section


def add_figure(doc: Document, image_path: Path, number: int, title: str, description: str) -> None:
    landscape_section(doc)
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.paragraph_format.space_after = Pt(7)
    run = heading.add_run(title)
    set_run_font(run, size=17, bold=True, color="000000")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(9.65))
    picture._inline.docPr.set("descr", description)
    picture._inline.docPr.set("title", title)
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f"Figura {number}. {description}")
    set_run_font(run, size=9, color=MUTED)
    portrait_section(doc)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    headers = rows[0]
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    set_table_borders(table)
    count = len(headers)
    total_width = 6.86
    if count == 3:
        if headers[0].strip().upper() in {"ID", "HALLAZGO"}:
            widths = [0.72, 4.18, 1.96] if headers[0].strip().upper() == "ID" else [1.45, 2.55, 2.86]
        else:
            widths = [1.8, 3.55, 1.95]
    elif count == 2:
        widths = [1.35, total_width - 1.35]
    else:
        widths = [total_width / count] * count

    all_rows = rows
    for row_index, values in enumerate(all_rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1.03
            if col_index == 0 and (headers[0].strip().upper() == "ID"):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(para, value.strip())
            for run in para.runs:
                if row_index == 0:
                    set_run_font(run, size=8.7, bold=True, color="FFFFFF")
                else:
                    set_run_font(run, size=8.5, color=INK)
        if row_index == 0:
            set_repeat_table_header(table.rows[0])
        set_row_cant_split(table.rows[row_index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run("Sistema de Préstamos EPCC  |  Primer examen  |  Página ")
        set_run_font(run, size=8, color=MUTED)
        add_field(paragraph, "PAGE")


def create_report() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    with tempfile.TemporaryDirectory(prefix="epcc-examen-") as temp_dir:
        temp = Path(temp_dir)
        domain_png = temp / "modelo-dominio.png"
        architecture_png = temp / "arquitectura-capas.png"
        draw_domain(domain_png)
        draw_architecture(architecture_png)
        figures = {"domain": domain_png, "architecture": architecture_png}

        lines = SOURCE.read_text(encoding="utf-8").splitlines()
        i = 0
        cover_title_added = False
        cover_subtitle = False
        cover_meta = False
        cover_open = True
        while i < len(lines):
            line = lines[i].rstrip()
            stripped = line.strip()
            if not stripped:
                i += 1
                continue
            if stripped == "{{PAGEBREAK}}":
                doc.add_page_break()
                cover_open = False
                cover_meta = False
                i += 1
                continue
            if stripped == "{{COVER_SUBTITLE}}":
                cover_subtitle = True
                i += 1
                continue
            if stripped == "{{COVER_META}}":
                cover_meta = True
                i += 1
                continue
            if stripped.startswith("{{FIGURE:"):
                key = stripped[len("{{FIGURE:") : -2]
                if key == "domain":
                    add_figure(doc, figures[key], 1, "Diagrama de clases del dominio", "Clases del esquema actual y multiplicidades de sus asociaciones.")
                elif key == "architecture":
                    add_figure(doc, figures[key], 2, "Diagrama de paquetes y clases de la arquitectura", "Paquetes de presentación, aplicación, dominio e infraestructura, con sus dependencias.")
                i += 1
                continue
            if stripped.startswith("|"):
                table_rows: list[list[str]] = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                        table_rows.append(cells)
                    i += 1
                add_table(doc, table_rows)
                continue
            heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
            if heading_match:
                level = len(heading_match.group(1))
                heading_text = heading_match.group(2)
                if level == 1 and not cover_title_added:
                    paragraph = doc.add_paragraph(style="Title")
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    ppr = paragraph._p.get_or_add_pPr()
                    border = ppr.find(qn("w:pBdr"))
                    if border is not None:
                        ppr.remove(border)
                    add_inline(paragraph, heading_text)
                    cover_title_added = True
                else:
                    paragraph = doc.add_paragraph(style=f"Heading {min(level, 3)}")
                    add_inline(paragraph, heading_text)
                i += 1
                continue
            if stripped.startswith("- ") or re.match(r"^\d+\.\s+", stripped):
                bullet = stripped.startswith("- ")
                body = stripped[2:] if bullet else re.sub(r"^\d+\.\s+", "", stripped)
                style_name = "List Bullet" if bullet else "List Number"
                paragraph = doc.add_paragraph(style=style_name)
                paragraph.paragraph_format.space_after = Pt(3)
                add_inline(paragraph, body)
                i += 1
                continue

            if cover_subtitle and not cover_meta:
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(26)
                add_inline(paragraph, stripped)
                for run in paragraph.runs:
                    set_run_font(run, size=16, color="000000")
                cover_subtitle = False
                i += 1
                continue
            if cover_meta:
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(6)
                add_inline(paragraph, stripped)
                for run in paragraph.runs:
                    set_run_font(run, size=11, color=MUTED)
                i += 1
                continue

            if cover_open:
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(8)
                add_inline(paragraph, stripped)
                for run in paragraph.runs:
                    set_run_font(run, size=12, color=INK)
                i += 1
                continue

            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.widow_control = True
            add_inline(paragraph, stripped)
            i += 1

    add_footer(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    create_report()
